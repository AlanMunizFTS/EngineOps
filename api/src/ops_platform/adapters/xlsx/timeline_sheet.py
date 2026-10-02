"""Builds the "Timeline" sheet: an activity table (mirroring the web
Schedule view's columns) plus a day-by-day Gantt grid to its right, styled
after the reference Excel template this replaces - dark-navy week/header
bands, today highlighted, weekends shaded, and bar colors matching the app's
own on-track/late/early-finish palette (see styles.py / SchedulePage.tsx)."""

from __future__ import annotations

from datetime import date, timedelta
from uuid import UUID

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ops_platform.adapters.xlsx.styles import (
    HEADER_FILL,
    HEADER_FONT,
    SEGMENT_FILLS,
    TODAY_FILL,
    WEEKEND_FILL,
    style_header_row,
)
from ops_platform.adapters.xlsx.timeline_rows import TimelineRow, order_timeline_rows
from ops_platform.domain.entities import Task
from ops_platform.domain.scheduling import (
    calculate_days_planned,
    calculate_schedule_status,
    compute_bar_segments,
    to_business_date,
    today_in_business_timezone,
    working_days_taken,
)

_INFO_COLUMNS = [
    "#",
    "Activity",
    "Responsible",
    "Start Date",
    "Days Planned",
    "Due Date",
    "Days Taken",
    "Status",
    "Close Date",
]
_INFO_WIDTHS = [4, 34, 16, 12, 12, 12, 12, 11, 12]
_WEEK_ROW = 1
_DOW_ROW = 2
_DAYNUM_ROW = 3
_DATA_START_ROW = 4
_DAY_COLUMN_WIDTH = 3.0


def write_timeline_sheet(
    ws: Worksheet, project_name: str, tasks: list[Task], member_names: dict[UUID, str]
) -> None:
    today = today_in_business_timezone()
    day_range = _build_day_range(tasks, today)
    day_start_col = len(_INFO_COLUMNS) + 1

    _write_info_header(ws, project_name)
    _write_day_headers(ws, day_range, today, day_start_col)

    rows = order_timeline_rows(tasks)
    for offset, row in enumerate(rows):
        row_index = _DATA_START_ROW + offset
        _write_info_cells(ws, row_index, offset + 1, row, member_names, today)
        _write_gantt_bar(ws, row_index, day_range, row.task, today, day_start_col)

    _apply_layout(ws, day_range, day_start_col)


def _build_day_range(tasks: list[Task], today: date) -> list[date]:
    candidates = [today]
    for task in tasks:
        if task.start_date:
            candidates.append(task.start_date)
        if task.due_date:
            candidates.append(task.due_date)
        if task.closed_at is not None:
            candidates.append(to_business_date(task.closed_at))

    start = min(candidates)
    start -= timedelta(days=start.weekday())
    end = max(candidates)
    end += timedelta(days=6 - end.weekday())

    days = []
    current = start
    while current <= end:
        days.append(current)
        current += timedelta(days=1)
    return days


def _write_info_header(ws: Worksheet, project_name: str) -> None:
    last_info_col = len(_INFO_COLUMNS)
    ws.merge_cells(start_row=_WEEK_ROW, start_column=1, end_row=_DOW_ROW, end_column=last_info_col)
    title_cell = ws.cell(row=_WEEK_ROW, column=1, value=f"{project_name} - Timeline")
    title_cell.font = Font(color="FFFFFF", bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in (_WEEK_ROW, _DOW_ROW):
        for col in range(1, last_info_col + 1):
            ws.cell(row=row, column=col).fill = HEADER_FILL

    for col, title in enumerate(_INFO_COLUMNS, start=1):
        ws.cell(row=_DAYNUM_ROW, column=col, value=title)
    style_header_row(ws, row=_DAYNUM_ROW, first_col=1, last_col=last_info_col)


def _write_day_headers(
    ws: Worksheet, day_range: list[date], today: date, day_start_col: int
) -> None:
    col = day_start_col
    index = 0
    while index < len(day_range):
        week_start_col = col
        week_num = day_range[index].isocalendar()[1]
        for _ in range(7):
            if index >= len(day_range):
                break
            _write_day_column(ws, col, day_range[index], today)
            col += 1
            index += 1
        _write_week_band(ws, week_start_col, col - 1, week_num)


def _write_week_band(ws: Worksheet, start_col: int, end_col: int, week_num: int) -> None:
    if end_col < start_col:
        return
    ws.merge_cells(
        start_row=_WEEK_ROW, start_column=start_col, end_row=_WEEK_ROW, end_column=end_col
    )
    cell = ws.cell(row=_WEEK_ROW, column=start_col, value=f"W{week_num}")
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center")
    for col in range(start_col, end_col + 1):
        ws.cell(row=_WEEK_ROW, column=col).fill = HEADER_FILL


def _write_day_column(ws: Worksheet, col: int, d: date, today: date) -> None:
    dow_cell = ws.cell(row=_DOW_ROW, column=col, value="SMTWTFS"[d.isoweekday() % 7])
    dow_cell.font = HEADER_FONT
    dow_cell.fill = HEADER_FILL
    dow_cell.alignment = Alignment(horizontal="center")

    daynum_cell = ws.cell(row=_DAYNUM_ROW, column=col, value=d.day)
    daynum_cell.alignment = Alignment(horizontal="center")
    if d == today:
        daynum_cell.fill = TODAY_FILL
        daynum_cell.font = Font(bold=True)
    elif d.weekday() >= 5:
        daynum_cell.fill = WEEKEND_FILL


def _write_info_cells(
    ws: Worksheet,
    row_index: int,
    number: int,
    row: TimelineRow,
    member_names: dict[UUID, str],
    today: date,
) -> None:
    task = row.task
    closed_at_date = to_business_date(task.closed_at) if task.closed_at else None
    schedule_status = calculate_schedule_status(today, task.due_date, closed_at_date)
    days_planned = calculate_days_planned(task.start_date, task.due_date)
    days_taken = (
        working_days_taken(task.start_date, closed_at_date or today)
        if task.start_date is not None
        else None
    )
    assignee = ", ".join(member_names.get(user_id, "") for user_id in task.assignee_ids)
    indent = "    " * row.depth
    values = [
        number,
        f"{indent}{task.title}",
        assignee,
        task.start_date.isoformat() if task.start_date else "",
        days_planned if days_planned is not None else "",
        task.due_date.isoformat() if task.due_date else "",
        days_taken if days_taken is not None else "",
        schedule_status.value.replace("_", " ").upper() if schedule_status else "",
        closed_at_date.isoformat() if closed_at_date else "",
    ]
    for col, value in enumerate(values, start=1):
        ws.cell(row=row_index, column=col, value=value)


def _write_gantt_bar(
    ws: Worksheet,
    row_index: int,
    day_range: list[date],
    task: Task,
    today: date,
    day_start_col: int,
) -> None:
    closed_at_date = to_business_date(task.closed_at) if task.closed_at else None
    schedule_status = calculate_schedule_status(today, task.due_date, closed_at_date)
    segments = compute_bar_segments(
        today, task.start_date, task.due_date, closed_at_date, schedule_status
    )
    day_index = {d: i for i, d in enumerate(day_range)}
    for segment in segments:
        fill = SEGMENT_FILLS[segment.kind]
        start_col = _col_for_date(segment.start, day_index, day_start_col, day_range)
        end_col = _col_for_date(segment.end, day_index, day_start_col, day_range)
        for col in range(start_col, end_col + 1):
            ws.cell(row=row_index, column=col).fill = fill


def _col_for_date(
    d: date, day_index: dict[date, int], day_start_col: int, day_range: list[date]
) -> int:
    """Clips to the grid's edges - a segment date can fall just outside
    `day_range` (e.g. `due_date + 1 day` landing the day after the range's
    last, week-aligned day) without it being worth widening the grid for."""
    if d in day_index:
        return day_start_col + day_index[d]
    if d < day_range[0]:
        return day_start_col
    return day_start_col + len(day_range) - 1


def _apply_layout(ws: Worksheet, day_range: list[date], day_start_col: int) -> None:
    for index, width in enumerate(_INFO_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for offset in range(len(day_range)):
        ws.column_dimensions[get_column_letter(day_start_col + offset)].width = _DAY_COLUMN_WIDTH
    ws.row_dimensions[_DAYNUM_ROW].height = 16
    ws.freeze_panes = ws.cell(row=_DATA_START_ROW, column=day_start_col).coordinate
