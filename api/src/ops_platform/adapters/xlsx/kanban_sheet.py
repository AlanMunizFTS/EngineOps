"""Builds the "Kanban" sheet: a flat, sortable/filterable table of every
board card (one row per Task), grouped by the Kanban column its status
currently maps to - a spreadsheet can't reproduce draggable swim-lanes, so a
table ordered by column is the closest useful equivalent."""

from __future__ import annotations

from uuid import UUID

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ops_platform.adapters.xlsx.styles import PRIORITY_FILLS, style_header_row
from ops_platform.domain.entities import KanbanBoard, Task, TaskStatus

_COLUMNS = ["Kanban Column", "Activity", "Responsible", "Priority", "Status", "Due Date"]
_COLUMN_WIDTHS = [20, 36, 20, 12, 14, 12]


def write_kanban_sheet(
    ws: Worksheet,
    tasks: list[Task],
    board: KanbanBoard | None,
    member_names: dict[UUID, str],
) -> None:
    ws.append(_COLUMNS)
    style_header_row(ws, row=1, first_col=1, last_col=len(_COLUMNS))

    board_tasks = [task for task in tasks if task.is_top_level]
    status_to_column = _status_to_column_name(board)
    column_order = _column_order(board)

    for task in sorted(
        board_tasks, key=lambda item: _sort_key(item, status_to_column, column_order)
    ):
        _write_row(ws, task, status_to_column.get(task.status, "Unassigned"), member_names)

    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    ws.freeze_panes = "A2"


def _status_to_column_name(board: KanbanBoard | None) -> dict[TaskStatus, str]:
    if board is None:
        return {}
    mapping: dict[TaskStatus, str] = {}
    for column in sorted(board.columns, key=lambda c: c.order_index):
        for status in column.maps_to_statuses:
            mapping.setdefault(status, column.name)
    return mapping


def _column_order(board: KanbanBoard | None) -> dict[str, int]:
    if board is None:
        return {}
    ordered = sorted(board.columns, key=lambda c: c.order_index)
    return {column.name: index for index, column in enumerate(ordered)}


def _sort_key(
    task: Task, status_to_column: dict[TaskStatus, str], column_order: dict[str, int]
) -> tuple[int, str]:
    column_name = status_to_column.get(task.status, "Unassigned")
    return (column_order.get(column_name, len(column_order)), task.title.lower())


def _write_row(ws: Worksheet, task: Task, column_name: str, member_names: dict[UUID, str]) -> None:
    assignee = member_names.get(task.assignee_id, "") if task.assignee_id else ""
    ws.append(
        [
            column_name,
            task.title,
            assignee,
            task.priority.value.capitalize(),
            task.status.value.replace("_", " ").title(),
            task.due_date.isoformat() if task.due_date else "",
        ]
    )
    priority_fill = PRIORITY_FILLS.get(task.priority.value)
    if priority_fill:
        cell = ws.cell(row=ws.max_row, column=4)
        cell.fill = PatternFill("solid", fgColor=priority_fill)
        cell.font = Font(color="FFFFFF", bold=True)
