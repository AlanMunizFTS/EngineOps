import io
import uuid
from datetime import date, timedelta

import openpyxl
from fastapi.testclient import TestClient


def _create_project(client: TestClient, auth_headers: dict[str, str]) -> str:
    response = client.post("/projects", json={"name": "Endforms Standard"}, headers=auth_headers)
    return response.json()["id"]


def test_export_404_for_unknown_project(client: TestClient, auth_headers: dict[str, str]) -> None:
    response = client.get(f"/projects/{uuid.uuid4()}/export.xlsx", headers=auth_headers)
    assert response.status_code == 404


def test_export_returns_workbook_with_kanban_and_timeline_sheets(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    project_id = _create_project(client, auth_headers)
    today = date.today()
    client.post(
        f"/projects/{project_id}/tasks",
        json={
            "title": "Cut fixture rework",
            "start_date": today.isoformat(),
            "due_date": (today + timedelta(days=3)).isoformat(),
        },
        headers=auth_headers,
    )

    response = client.get(f"/projects/{project_id}/export.xlsx", headers=auth_headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Kanban", "Timeline"]

    kanban_ws = workbook["Kanban"]
    assert [cell.value for cell in kanban_ws[1]] == [
        "Kanban Column",
        "Activity",
        "Responsible",
        "Priority",
        "Status",
        "Due Date",
    ]
    kanban_titles = [row[1].value for row in kanban_ws.iter_rows(min_row=2)]
    assert "Cut fixture rework" in kanban_titles

    timeline_ws = workbook["Timeline"]
    timeline_activity_cells = [row[1].value for row in timeline_ws.iter_rows(min_row=4)]
    assert any(
        value is not None and "Cut fixture rework" in value for value in timeline_activity_cells
    )


def test_export_includes_only_top_level_tasks_in_kanban_sheet(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    project_id = _create_project(client, auth_headers)
    parent = client.post(
        f"/projects/{project_id}/tasks", json={"title": "Parent activity"}, headers=auth_headers
    ).json()
    client.post(
        f"/tasks/{parent['id']}/subtasks",
        json={"title": "Child activity"},
        headers=auth_headers,
    )

    response = client.get(f"/projects/{project_id}/export.xlsx", headers=auth_headers)
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    kanban_titles = {row[1].value for row in workbook["Kanban"].iter_rows(min_row=2)}

    assert "Parent activity" in kanban_titles
    assert "Child activity" not in kanban_titles
